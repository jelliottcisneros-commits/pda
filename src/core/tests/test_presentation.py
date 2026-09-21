import io
from types import SimpleNamespace
from unittest import TestCase

from pptx import Presentation

from core.presentation import generate_group_presentation


def _attach_default_area_scores(scores):
    """Give synthetic scores the seven sociocultural area objects."""
    area_fields = (
        "Religion_Score",
        "Disability_Score",
        "Culture_Score",
        "Gender_Score",
        "Race_Score",
        "Class_Score",
        "Sexual_Orientation_Score",
    )

    for score in scores:
        for field in area_fields:
            setattr(
                score,
                field,
                SimpleNamespace(
                    sensitivity=0,
                    oneness=0,
                    strength=0,
                    appreciation=0,
                    leveraged=8,
                ),
            )


class GroupPresentationTests(TestCase):
    def test_generate_group_presentation_creates_editable_pptx(self):
        scores = [
            SimpleNamespace(
                sensitivity_total=50,
                oneness_total=30,
                strength_total=20,
                appreciation_total=30,
                leveraged_total=40,
            ),
            SimpleNamespace(
                sensitivity_total=20,
                oneness_total=50,
                strength_total=30,
                appreciation_total=30,
                leveraged_total=40,
            ),
            SimpleNamespace(
                sensitivity_total=20,
                oneness_total=30,
                strength_total=50,
                appreciation_total=30,
                leveraged_total=40,
            ),
        ]

        _attach_default_area_scores(scores)
        result = generate_group_presentation(scores)

        self.assertIsInstance(result, io.BytesIO)

        result.seek(0)
        prs = Presentation(result)

        self.assertGreaterEqual(len(prs.slides), 1)

        slide_text = "\n".join(
            shape.text
            for shape in prs.slides[0].shapes
            if hasattr(shape, "text")
        )

        self.assertIn("Primary Pattern Distribution", slide_text)
        self.assertIn("Sensitivity", slide_text)
        self.assertIn("Oneness", slide_text)
        self.assertIn("Strength", slide_text)

        # This fixture has one clear primary-pattern person in each category.
        self.assertIn("1", slide_text)

    def test_generate_group_presentation_includes_group_anya_slide(self):
        scores = [
            SimpleNamespace(
                sensitivity_total=10,
                oneness_total=20,
                strength_total=30,
                appreciation_total=40,
                leveraged_total=50,
            ),
            SimpleNamespace(
                sensitivity_total=20,
                oneness_total=30,
                strength_total=40,
                appreciation_total=50,
                leveraged_total=40,
            ),
        ]

        _attach_default_area_scores(scores)
        result = generate_group_presentation(scores)

        result.seek(0)
        prs = Presentation(result)

        self.assertGreaterEqual(len(prs.slides), 2)

        slide_text = "\n".join(
            shape.text
            for shape in prs.slides[1].shapes
            if hasattr(shape, "text")
        )

        self.assertIn("A/NYA Power Quotient", slide_text)
        self.assertIn("Actualized Power", slide_text)
        self.assertIn("Not-Yet-Actualized Power", slide_text)
        self.assertIn("53.2%", slide_text)
        self.assertIn("46.8%", slide_text)


    def test_generate_group_presentation_includes_area_insight_slide(self):
        scores = [
            SimpleNamespace(
                sensitivity_total=30,
                oneness_total=20,
                strength_total=10,
                appreciation_total=20,
                leveraged_total=40,
            ),
            SimpleNamespace(
                sensitivity_total=20,
                oneness_total=30,
                strength_total=10,
                appreciation_total=20,
                leveraged_total=40,
            ),
        ]

        _attach_default_area_scores(scores)

        # Person 1: Religion has the greatest NYA.
        scores[0].Religion_Score = SimpleNamespace(
            sensitivity=2,
            oneness=4,
            strength=6,
            appreciation=8,
            leveraged=0,
        )

        # Person 2: Disability has the greatest NYA.
        scores[1].Disability_Score = SimpleNamespace(
            sensitivity=2,
            oneness=4,
            strength=6,
            appreciation=8,
            leveraged=0,
        )

        result = generate_group_presentation(scores)

        result.seek(0)
        prs = Presentation(result)

        self.assertGreaterEqual(len(prs.slides), 3)

        slide_text = "\n".join(
            shape.text
            for shape in prs.slides[2].shapes
            if hasattr(shape, "text")
        )

        self.assertIn("A/NYA by Sociocultural Area", slide_text)
        self.assertIn("Most Room for Growth", slide_text)
        self.assertIn("Strongest Current Asset", slide_text)
        self.assertIn("Religion", slide_text)
        self.assertIn("Disability", slide_text)

    def test_primary_pattern_slide_includes_group_pattern_and_learning_edges(self):
        scores = [
            SimpleNamespace(
                sensitivity_total=50,
                oneness_total=20,
                strength_total=10,
                appreciation_total=20,
                leveraged_total=40,
            ),
            SimpleNamespace(
                sensitivity_total=40,
                oneness_total=20,
                strength_total=10,
                appreciation_total=20,
                leveraged_total=40,
            ),
            SimpleNamespace(
                sensitivity_total=10,
                oneness_total=50,
                strength_total=20,
                appreciation_total=20,
                leveraged_total=40,
            ),
        ]

        _attach_default_area_scores(scores)

        result = generate_group_presentation(scores)

        result.seek(0)
        prs = Presentation(result)

        slide_text = "\n".join(
            shape.text
            for shape in prs.slides[0].shapes
            if hasattr(shape, "text")
        )

        self.assertIn("Group Primary Pattern", slide_text)
        self.assertIn("Learning Edges", slide_text)

        # Group totals:
        # Sensitivity = 100, Oneness = 90, Strength = 40
        self.assertIn("Group Primary Pattern: Sensitivity", slide_text)

        # Learning-edge mapping:
        # 2 Sensitivity primaries -> Strength edge
        # 1 Oneness primary -> Sensitivity edge
        self.assertIn("Strength: 2", slide_text)
        self.assertIn("Sensitivity: 1", slide_text)
        self.assertIn("Oneness: 0", slide_text)

    def test_generate_group_presentation_includes_appreciation_slide(self):
        scores = [
            SimpleNamespace(
                sensitivity_total=30,
                oneness_total=20,
                strength_total=10,
                appreciation_total=40,
                leveraged_total=40,
            ),
            SimpleNamespace(
                sensitivity_total=20,
                oneness_total=30,
                strength_total=10,
                appreciation_total=50,
                leveraged_total=40,
            ),
            SimpleNamespace(
                sensitivity_total=20,
                oneness_total=10,
                strength_total=30,
                appreciation_total=30,
                leveraged_total=40,
            ),
        ]

        _attach_default_area_scores(scores)

        # Appreciation values across selected areas.
        scores[0].Religion_Score.appreciation = 5
        scores[1].Religion_Score.appreciation = 2
        scores[2].Religion_Score.appreciation = 4

        scores[0].Race_Score.appreciation = 6
        scores[1].Race_Score.appreciation = 1
        scores[2].Race_Score.appreciation = 4

        result = generate_group_presentation(scores)

        result.seek(0)
        prs = Presentation(result)

        self.assertGreaterEqual(len(prs.slides), 4)

        slide_text = "\n".join(
            shape.text
            for shape in prs.slides[3].shapes
            if hasattr(shape, "text")
        )

        self.assertIn("Appreciation", slide_text)
        self.assertIn("120 out of 168 possible", slide_text)
        self.assertIn("Religion", slide_text)
        self.assertIn("Race", slide_text)
        self.assertIn("4 or more", slide_text)

    def test_generate_group_presentation_includes_appreciation_slide(self):
        scores = [
            SimpleNamespace(
                sensitivity_total=30,
                oneness_total=20,
                strength_total=10,
                appreciation_total=40,
                leveraged_total=40,
            ),
            SimpleNamespace(
                sensitivity_total=20,
                oneness_total=30,
                strength_total=10,
                appreciation_total=50,
                leveraged_total=40,
            ),
            SimpleNamespace(
                sensitivity_total=20,
                oneness_total=10,
                strength_total=30,
                appreciation_total=30,
                leveraged_total=40,
            ),
        ]

        _attach_default_area_scores(scores)

        # Appreciation values across selected areas.
        scores[0].Religion_Score.appreciation = 5
        scores[1].Religion_Score.appreciation = 2
        scores[2].Religion_Score.appreciation = 4

        scores[0].Race_Score.appreciation = 6
        scores[1].Race_Score.appreciation = 1
        scores[2].Race_Score.appreciation = 4

        result = generate_group_presentation(scores)

        result.seek(0)
        prs = Presentation(result)

        self.assertGreaterEqual(len(prs.slides), 4)

        slide_text = "\n".join(
            shape.text
            for shape in prs.slides[3].shapes
            if hasattr(shape, "text")
        )

        self.assertIn("Appreciation", slide_text)
        self.assertIn("120 out of 168 possible", slide_text)
        self.assertIn("Religion", slide_text)
        self.assertIn("Race", slide_text)
        self.assertIn("4 or more", slide_text)

    def test_primary_pattern_slide_includes_group_pattern_totals(self):
        scores = [
            SimpleNamespace(
                sensitivity_total=50,
                oneness_total=20,
                strength_total=10,
                appreciation_total=20,
                leveraged_total=40,
            ),
            SimpleNamespace(
                sensitivity_total=40,
                oneness_total=30,
                strength_total=20,
                appreciation_total=20,
                leveraged_total=40,
            ),
            SimpleNamespace(
                sensitivity_total=10,
                oneness_total=40,
                strength_total=30,
                appreciation_total=20,
                leveraged_total=40,
            ),
        ]

        _attach_default_area_scores(scores)

        result = generate_group_presentation(scores)

        result.seek(0)
        prs = Presentation(result)

        slide_text = "\n".join(
            shape.text
            for shape in prs.slides[0].shapes
            if hasattr(shape, "text")
        )

        self.assertIn("Group Pattern Scores", slide_text)
        self.assertIn("Sensitivity: 100", slide_text)
        self.assertIn("Oneness: 90", slide_text)
        self.assertIn("Strength: 60", slide_text)

    def test_generate_group_presentation_includes_individual_anya_distribution(self):
        scores = [
            SimpleNamespace(
                sensitivity_total=10,
                oneness_total=20,
                strength_total=30,
                appreciation_total=40,
                leveraged_total=50,
            ),
            SimpleNamespace(
                sensitivity_total=20,
                oneness_total=30,
                strength_total=40,
                appreciation_total=50,
                leveraged_total=40,
            ),
            SimpleNamespace(
                sensitivity_total=30,
                oneness_total=20,
                strength_total=10,
                appreciation_total=20,
                leveraged_total=30,
            ),
        ]

        _attach_default_area_scores(scores)

        result = generate_group_presentation(scores)

        result.seek(0)
        prs = Presentation(result)

        self.assertGreaterEqual(len(prs.slides), 5)

        slide_text = "\n".join(
            shape.text
            for shape in prs.slides[4].shapes
            if hasattr(shape, "text")
        )

        self.assertIn("Individual A/NYA Distribution", slide_text)
        self.assertIn("Actualized Power", slide_text)
        self.assertIn("Not-Yet-Actualized Power", slide_text)
        self.assertIn("Each point represents one respondent", slide_text)

    def test_individual_anya_axis_runs_from_100_actualized_to_0_actualized(self):
        scores = [
            SimpleNamespace(
                sensitivity_total=10,
                oneness_total=20,
                strength_total=30,
                appreciation_total=40,
                leveraged_total=50,
            ),
        ]

        _attach_default_area_scores(scores)

        result = generate_group_presentation(scores)
        result.seek(0)
        prs = Presentation(result)

        slide = prs.slides[4]

        percentage_shapes = {
            shape.text: shape.left
            for shape in slide.shapes
            if hasattr(shape, "text") and shape.text in {"100%", "0%"}
        }

        self.assertIn("100%", percentage_shapes)
        self.assertIn("0%", percentage_shapes)

        # 100% Actualized must be physically LEFT of 0% Actualized.
        self.assertLess(
            percentage_shapes["100%"],
            percentage_shapes["0%"],
        )
